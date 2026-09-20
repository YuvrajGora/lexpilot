import io
from fastapi import HTTPException, status

def extract_text_from_txt(file_bytes: bytes) -> str:
    """Safely extracts text from TXT using UTF-8, falling back to latin-1."""
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        try:
            return file_bytes.decode("latin-1")
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Unable to decode TXT file. Please ensure it is text-encoded."
            )

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts text from PDF pages using pypdf."""
    try:
        from pypdf import PdfReader
        stream = io.BytesIO(file_bytes)
        reader = PdfReader(stream)
        
        extracted_pages = []
        for idx, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                extracted_pages.append(text)
                
        return "\n\n".join(extracted_pages)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="The document could not be read or does not contain extractable text."
        )

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts text from DOCX paragraphs and tables using python-docx."""
    try:
        import docx
        stream = io.BytesIO(file_bytes)
        doc = docx.Document(stream)
        
        extracted = []
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                extracted.append(text)
                
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    extracted.append(" | ".join(row_texts))
                    
        return "\n\n".join(extracted)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="The document could not be read or does not contain extractable text."
        )

def extract_document_text(file_bytes: bytes, extension: str) -> str:
    """Dispatches text extraction based on file extension."""
    ext = extension.lower()
    if ext == ".txt":
        return extract_text_from_txt(file_bytes)
    elif ext == ".pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext == ".docx":
        return extract_text_from_docx(file_bytes)
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file format: {ext}"
        )
